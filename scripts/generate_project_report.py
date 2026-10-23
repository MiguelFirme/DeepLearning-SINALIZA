"""Gera o relatório histórico do Sinaliza a partir dos registros auditados.

Atualize as afirmações e as fontes abaixo quando novos experimentos forem
concluídos. Execute com o Python de documentos do ambiente Codex.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Relatorio_historico_Sinaliza_2026-10-22.docx"
DARK = RGBColor(29, 50, 76)
PALE = "EDF2F7"
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(85, 93, 104)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:color"), "D9D9D9")
        el.set(qn("w:sz"), "4")
        item.append(el)
    tc_pr.append(item)


def cell_margin(cell, top=95, start=110, bottom=95, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, size in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(size))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (head, width) in enumerate(zip(headers, widths)):
        cell = t.rows[0].cells[i]
        cell.width = Cm(width)
        cell.text = head
        shade(cell, "1D324C")
        cell_margin(cell)
        borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(8.3)
                run.font.color.rgb = WHITE
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for ri, record in enumerate(rows):
        row = t.add_row()
        for i, (value, width) in enumerate(zip(record, widths)):
            cell = row.cells[i]
            cell.width = Cm(width)
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)
            borders(cell)
            if ri % 2:
                shade(cell, PALE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                if i and len(record) >= 3 and len(str(value)) < 22:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def para(doc, text, *, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def heading(doc, title, level=1):
    p = doc.add_heading(title, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    return p


def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Sinaliza  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.75)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Title", 20, 0, 10), ("Heading 1", 13.2, 14, 5), ("Heading 2", 10.8, 9, 3)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    doc.styles["List Bullet"].font.name = "Aptos"
    doc.styles["List Bullet"].font.size = Pt(9.8)

    title = doc.add_paragraph("Relatório histórico do projeto Sinaliza", style="Title")
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    p = para(doc, "Desenvolvimento, testes e resultados até 22 de outubro de 2026")
    for r in p.runs:
        r.font.color.rgb = GRAY
        r.font.size = Pt(10)
    para(doc, "Projeto de extensão em Libras e visão computacional. Este relatório registra os experimentos concluídos e o estado verificável do protótipo para apoiar a continuidade do trabalho pela equipe e pelo professor.")

    heading(doc, "Resumo executivo")
    para(doc, "O Sinaliza hoje demonstra classificação de sinais isolados de Libras por webcam. A aplicação React e FastAPI carrega o Transformer de oito classes com pose e face, cujo checkpoint acertou 40/40 vídeos de validação e 37/40 de teste exploratório no MINDS-Libras. Esse resultado não mede tradução de frases nem desempenho confiável de usuários novos. [R1, R8]")
    para(doc, "O experimento mais recente para um vocabulário maior combina dois Transformers de 15 classes. A fusão top 5 elevou a validação de 49/75 para 52/75 acertos e o teste exploratório de 41/75 para 44/75. A inversão da ordem dos modelos não mudou as previsões top 5 nesses vídeos. A fusão foi testada offline; não está integrada ao backend nem teve sua latência medida. [R10, R11]")
    para(doc, "A principal limitação ao longo do projeto é a generalização entre pessoas. No V-Librasil havia aproximadamente um vídeo por classe e articulador; no MINDS-Libras, a cobertura de mãos difere muito entre os sinalizadores de treino e os retidos. O teste do Sinalizador12 já foi consultado em vários experimentos e deve ser tratado como exploratório. [R2, R5, R7]")

    heading(doc, "1 Objetivo e protocolo de avaliação")
    para(doc, "A meta de longo prazo é traduzir Libras para texto natural em contextos escolares, domésticos e de recepção. Os dados e modelos atuais reconhecem palavras isoladas; não há corpus de vídeos de frases com transcrições para treinar tradução contínua. Por isso, as métricas deste relatório são de classificação top 1 de um sinal entre classes conhecidas. [R1, R6]")
    table(doc, ["Conjunto", "Escala", "Divisão usada", "Observação"], [
        ("V-Librasil", "4.086 vídeos; 1.364 classes", "Dois articuladores no treino; terceiro no teste", "Geralmente só dois vídeos reais de treino por classe."),
        ("MINDS-Libras", "800 vídeos; 20 classes; 8 sinalizadores", "Seis sinalizadores no treino, 11 na validação, 12 no teste", "40 vídeos por classe; 30/5/5 no protocolo por classe."),
    ], [3.0, 3.4, 5.2, 5.2])
    para(doc, "Os splits do MINDS-Libras não compartilham sinalizadores. O vídeo inteiro e todos os seus quadros ficam no mesmo split. Augmentation dinâmica é aplicada apenas no treino; validação e teste usam landmarks originais. O checkpoint é selecionado pela menor perda de validação, exceto nas experiências exploratórias com grade explicitamente identificadas adiante. [R3, R6]")

    heading(doc, "2 Representação e treinamento")
    para(doc, "O MediaPipe Holistic extrai 21 pontos por mão, 25 pontos de pose selecionados e 40 de face, totalizando 346 valores brutos por quadro. A normalização usa a referência dos ombros quando disponível; são acrescentadas velocidades e cinco distâncias, resultando em 697 features por quadro. Cada vídeo é amostrado para 48 quadros. No modo pose e face, os 42 pontos detalhados das mãos são zerados, mas ombros, cotovelos, pulsos e face permanecem. [R4]")
    para(doc, "Foram implementados BiLSTM, Transformer temporal 1D e TCN; os experimentos comparativos mais completos desta fase usaram BiLSTM e Transformer. A configuração Transformer do MINDS-Libras emprega dimensão 64, duas camadas, quatro cabeças, feedforward 128 e dropout 0,45. O pipeline inclui Adam com weight decay de 1e-4, ReduceLROnPlateau, early stopping por val_loss e pequenas rotações, escalas e ruído apenas no treino. Os parâmetros exatos ficam nos YAMLs e checkpoints arquivados. [R3, R8, R10]")
    para(doc, "A seleção de vocabulário para o V-Librasil foi incorporada ao código: annotations.csv é lido na preparação, verbos e termos essenciais são protegidos por regras e exceções editáveis, e o processo grava classes aceitas e excluídas. Essa proteção depende de revisão humana em casos ambíguos. [R2]")

    heading(doc, "3 Etapa V Librasil")
    para(doc, "Os primeiros testes de BiLSTM com 50 classes deram 12%, 12% e 8% de top 1 nos três articuladores retidos. Os 28% observados numa rodada eram acurácia de treino, não teste. Com o vocabulário filtrado de cerca de 398 classes, uma rodada chegou a 97,86% no treino e apenas 3,78% no articulador reservado, evidenciando forte distância entre memorizar exemplos e reconhecer outra pessoa. [R2, R5]")
    table(doc, ["Experimento", "Teste por articulador", "Leitura"], [
        ("100 classes escolares e necessidades", "7%, 12%, 7%", "Um vídeo de teste por classe e pessoa; melhora relativa ao acaso de 1%, mas baixa utilidade prática."),
        ("23 classes da categoria escola", "1/23, 4/23, 1/23", "Treino próximo de 100%; generalização fraca."),
        ("23 classes com 7 variações sintéticas por classe", "2/23, 5/23, 3/23", "10/69 no total, contra 6/69 antes; derivados não são pessoas novas."),
        ("Grade de arquitetura e dropout", "3/23 no refit exploratório", "1 camada e 64 unidades igualou top 1 do modelo maior com menos parâmetros."),
        ("BiLSTM regularizada sem sintéticos", "1/23", "Menor memorização, sem ganho de generalização neste conjunto mínimo."),
    ], [5.4, 3.5, 7.9])
    para(doc, "A geração sintética preservou os 69 arquivos originais das 23 classes e criou 161 variações pequenas para chegar a dez arquivos por classe. O treino posterior usou derivados apenas dos articuladores de treino. Houve ganho de quatro acertos entre 69 vídeos reais de teste, mas as variações derivam das mesmas três execuções originais por classe. [R5]")
    para(doc, "Uma grade pequena examinou camadas, largura, taxa de aprendizado e dropout. A revisão escolheu 1 camada BiLSTM, 64 unidades por direção, learning rate 0,0003 e dropout 0,2 pela validação, mas o refit fez 3/23 no articulador de teste. Regularização mais forte com 23 vídeos reais de treino também não resolveu a falta de diversidade de pessoas. [R3]")

    heading(doc, "4 Migração para MINDS Libras")
    para(doc, "O MINDS-Libras foi adotado para aumentar repetições e número de pessoas. O piloto começou com cinco classes e 200 vídeos, depois foi ampliado para oito classes e 320 vídeos e, por fim, para as 20 classes e 800 vídeos. Os conjuntos de 15 classes reutilizam os 800 landmarks, mas selecionam 600 vídeos. A comparação deve respeitar o número de classes: uma acurácia maior em oito classes não equivale a melhor cobertura de um vocabulário de 15 ou 20. [R6, R8, R9, R10]")
    heading(doc, "Piloto de cinco classes", level=2)
    table(doc, ["Modelo e entrada", "Classes", "Validação", "Teste", "Notas"], [
        ("BiLSTM completa", "5", "5/25 (20%)", "3/25 (12%)", "Primeiro piloto; val estagnada no acaso."),
        ("BiLSTM pose e face", "5", "11/25 (44%)", "15/25 (60%)", "Melhora ao retirar mãos detalhadas."),
        ("I3D WLASL congelado", "5", "10/25 (40%)", "Não feito", "Encoder RGB + cabeça linear; treino 100%."),
        ("Transformer pose e face", "5", "19/25 (76%)", "20/25 (80%)", "Melhor no piloto de cinco classes."),
        ("Transformer completo", "5", "5/25 (20%)", "5/25 (20%)", "Previu Banheiro para todos os vídeos."),
    ], [5.0, 1.2, 2.7, 2.7, 5.2])
    heading(doc, "Expansão para oito classes", level=2)
    table(doc, ["Modelo e entrada", "Classes", "Validação", "Teste", "Notas"], [
        ("Transformer pose e face", "8", "40/40 (100%)", "37/40 (92,5%)", "Checkpoint ativo na aplicação."),
        ("Transformer só mãos", "8", "5/40 (12,5%)", "6/40 (15%)", "Ramo de mãos isolado fraco."),
        ("Transformer completo", "8", "23/40 (57,5%)", "35/40 (87,5%)", "Abaixo do pose e face no mesmo split."),
    ], [5.0, 1.2, 2.7, 2.7, 5.2])
    para(doc, "As oito classes são Acontecer, Aluno, Banheiro, Barulho, Conhecer, Medo, Ruim e Vontade. [R8]")
    para(doc, "No piloto de cinco classes, o I3D foi carregado a partir de pesos oficiais pré treinados em WLASL2000 e teve o encoder congelado. A cabeça linear aprendeu os 150 vídeos de treino, mas fez 10/25 na validação, ante 11/25 do BiLSTM pose e face e 19/25 do Transformer pose e face. Não foi feito fine tuning do encoder nem avaliação de teste desse I3D. [R6, R7]")

    heading(doc, "5 Cobertura e detecção das mãos")
    para(doc, "Na auditoria dos 320 vídeos de oito classes, ao menos uma mão foi detectada em 48,77% dos quadros de treino e 99,75% dos de validação. Ambas as mãos apareceram em 9,93% e 95,44%, respectivamente. Isso descreve a saída do MediaPipe, não a visibilidade real nem a qualidade geométrica dos pontos; inclui preparação e repouso. Para sinais predominantemente de uma mão, como Medo, exigir duas mãos seria uma regra inadequada. [R7]")
    para(doc, "Em 240 quadros amostrados de seis vídeos do piloto, o Holistic registrou pelo menos uma mão em 182 quadros e ambas em 137; o MediaPipe Hands separado fez 179 e 125. Recortes guiados pela pose fizeram 177 e 117. Nenhuma dessas duas alternativas recuperou de modo consistente as mãos ausentes. Não se substituiu o extrator nem se inventaram coordenadas faltantes. [R7]")

    heading(doc, "6 Ampliação para 20 e recorte de 15 classes")
    table(doc, ["Configuração", "Treino validação teste", "Validação", "Teste exploratório"], [
        ("Transformer pose e face, 20 classes", "600 / 100 / 100", "37/100 (37%)", "Não executado"),
        ("Transformer pose e face, 15 classes", "450 / 75 / 75", "49/75 (65,3%)", "41/75 (54,7%)"),
        ("Transformer pose face mãos, 15 classes", "450 / 75 / 75", "45/75 (60,0%)", "34/75 (45,3%)"),
        ("Fusão sequencial top 5, 15 classes", "Mesmos 75 vídeos por split", "52/75 (69,3%)", "44/75 (58,7%)"),
    ], [6.2, 3.5, 3.4, 3.7])
    para(doc, "No checkpoint de 20 classes, 11 classes tiveram 0/5 acertos na validação. A equipe retirou America, Esquina, Sapo, Bala e Maca e retreinou do zero com 15 classes; a tarefa e a camada de saída mudaram, portanto a melhora percentual não isola o efeito de uma única palavra. O modelo completo de 15 classes recuperou algumas classes que pose e face errava, mas caiu no agregado. [R9, R10]")
    para(doc, "O vocabulário de 15 classes é Acontecer, Aluno, Amarelo, Aproveitar, Banco, Banheiro, Barulho, Cinco, Conhecer, Espelho, Filho, Medo, Ruim, Vacina e Vontade. [R10]")
    para(doc, "A fusão avaliada em 21/10/2026 usa dois checkpoints já treinados. O primeiro restringe as classes ao top 5; o escore para cada candidata é 0,65 × P(pose e face) + 0,35 × P(completo). Na validação, corrigiu oito erros do modelo pose e face e introduziu cinco; no teste exploratório, corrigiu seis e introduziu três. Inverter a ordem produziu exatamente as mesmas 75 previsões top 5 em cada split. Top 3 alterou algumas decisões, mas não melhorou o resultado principal. [R11]")
    para(doc, "A fusão representa o melhor resultado observado entre as variantes de 15 classes, com ganho líquido de três vídeos em cada split. Seus escores softmax não foram calibrados e o método ainda não foi medido em tempo real. O sinalizador de teste já havia sido consultado em trabalhos anteriores, portanto o ganho é exploratório. [R11]")

    heading(doc, "7 Aplicação web e verificações").paragraph_format.page_break_before = True
    para(doc, "O frontend React e Vite envia quadros JPEG da webcam por WebSocket à API FastAPI. O backend extrai landmarks, mantém o buffer temporal, carrega o checkpoint de oito classes e devolve palavra, candidatos e tempos. Um teste de integração enviou 48 JPEGs e recebeu eventos de pronto, confirmação de quadro e predição. O build do frontend passou. Um vídeo de validação conhecido, Banheiro, foi reconhecido no fluxo web com aproximadamente 79,2% de confiança; não é teste de uma pessoa inédita. [R1, R12]")
    para(doc, "A primeira previsão exige 48 quadros processados. A mudança de latência reduziu o avanço posterior da janela de 24 para 12 quadros e passou a mostrar FPS de análise e tempo de ida e volta; não há medição consolidada do desempenho com a Canon R100 neste relatório. O tempo do Transformer isolado não inclui JPEG e MediaPipe. [R13]")
    para(doc, "Um filtro de movimento para cenas paradas foi testado e revertido porque restringiu demais as palavras reconhecidas na webcam. O limiar de exibição ativo voltou a 0,30; permanece a verificação de presença de pose e face. Como o modelo foi treinado apenas com classes positivas, pode atribuir uma palavra a uma pessoa parada ou a um gesto fora do vocabulário. [R1, R12]")
    para(doc, "Os checkpoints, mapas de classes, splits, configurações, históricos, métricas e logs foram arquivados em Modelos salvos, experiments e logs. O modelo de fusão não substituiu o checkpoint de oito classes configurado no backend. [R1, R10, R11]")

    heading(doc, "8 Conclusões e próximos testes")
    bullet(doc, "Recrutar sinalizadores e sessões novos, com câmeras e iluminação variadas. Congelar protocolo, pesos e regra de fusão antes de avaliar um teste realmente inédito.")
    bullet(doc, "Medir cobertura e qualidade da mão ativa durante o gesto, não apenas a presença de pontos ao longo do vídeo. Corrigir gravações cortadas, borradas ou ocluídas antes de mudar novamente o detector.")
    bullet(doc, "Medir latência ponta a ponta e taxa de acertos e falsos positivos na webcam com o modelo atual e, depois, com a fusão de 15 classes. A segunda inferência pode aumentar o tempo de resposta.")
    bullet(doc, "Gravar exemplos negativos como repouso, conversa e movimentos incompletos para avaliar rejeição de não sinais. Para tradução de frases, coletar vídeos de sentenças com transcrições e definir outra tarefa de treinamento.")
    para(doc, "O resultado mais sólido é metodológico: splits por pessoa, rastreio de origem de variantes, checkpoint por validação e registros por vídeo permitiram encontrar tanto ganhos quanto regressões. A melhor configuração offline de 15 classes ainda precisa ser confirmada em pessoas novas antes de ser integrada à interface.")

    heading(doc, "Fontes do projeto")
    sources = [
        ("R1", "README.md; backend/config.py"),
        ("R2", "docs/PLANO_FILTRAGEM_TREINO.md; docs/TRAINING.md"),
        ("R3", "docs/PLANO_LOGS_GRID_ESCOLA.md; docs/PLANO_REGULARIZACAO_SEM_VAZAMENTO.md"),
        ("R4", "docs/AUDITORIA_PONTOS_MEDIAPIPE.md; ml/features/landmarks.py"),
        ("R5", "docs/PLANO_ESCOLA_NECESSIDADES_100.md; docs/PLANO_TESTE_CATEGORIA_ESCOLA.md; docs/PLANO_VARIACOES_LANDMARKS_ESCOLA.md; docs/PLANO_AUMENTACAO_LANDMARKS.md"),
        ("R6", "docs/PLANO_MIGRACAO_MINDS_LIBRAS.md; docs/PLANO_TRANSFORMER_1D_MINDS.md"),
        ("R7", "docs/PLANO_TESTE_I3D_WLASL.md; docs/AUDITORIA_MAOS_MINDS_8_CLASSES.md; docs/PLANO_ABLACAO_MAOS_MINDS.md"),
        ("R8", "docs/PLANO_TRANSFORMER_MINDS_8_CLASSES.md; Modelos salvos/README.md"),
        ("R9", "docs/PLANO_TRANSFORMER_MINDS_20_CLASSES.md; experiments/transformer_minds_all20_pose_face_seed42/val_metrics.json"),
        ("R10", "docs/EXPERIMENTO_TRANSFORMER_MINDS_15_CLASSES_POSE_FACE_2026-10-21.md; docs/EXPERIMENTO_TRANSFORMER_15_CLASSES_FULL_2026-10-21.md"),
        ("R11", "docs/EXPERIMENTO_FUSAO_SEQUENCIAL_15_CLASSES_2026-10-21.md; experiments/fusion_15_classes/"),
        ("R12", "docs/PLANO_INTEGRACAO_WEB_8_CLASSES.md; docs/PLANO_REJEICAO_SEM_SINAL.md"),
        ("R13", "docs/PLANO_LATENCIA_WEBCAM_2026-10-14.md"),
    ]
    for code, paths in sources:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(code + "  ")
        r.bold = True
        r.font.size = Pt(7.4)
        r = p.add_run(paths)
        r.font.size = Pt(7.4)
    para(doc, "Frações de acertos são a referência principal; percentuais estão arredondados.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
